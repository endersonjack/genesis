from datetime import timedelta
from decimal import Decimal
from django.test import TestCase, RequestFactory
from django.utils import timezone
from django.template.loader import render_to_string
from .busca import BuscaPagamentosForm, termo_corresponde
from .views import _buscar_pagamentos_context, _filtros_export_busca_pagamentos
from .models import (Caixa, CategoriaFinanceira, ContaBancaria, PagamentoNotaFiscal,
    PagamentoNotaFiscalItem, BoletoPagamento, PagamentoImposto, PagamentoImpostoItem,
    AutoridadeTributaria, PagamentoBancarioAvulso)
from empresas.models import Empresa
from fornecedores.models import Fornecedor


class BuscaAvancadaTests(TestCase):
    def setUp(self):
        self.factory = RequestFactory()
        self.empresa = Empresa.objects.create(razao_social="Empresa busca", cnpj="12345678000199")
        self.caixa = Caixa.objects.get(empresa=self.empresa, tipo=Caixa.Tipo.GERAL)
        self.fornecedor = Fornecedor.objects.create(empresa=self.empresa, tipo="PJ",
            nome="Fornecedor Aurora", razao_social="Aurora Materiais", cpf_cnpj="12345678000195")
        self.today = timezone.localdate()
        self.nf = PagamentoNotaFiscal.objects.create(empresa=self.empresa, caixa=self.caixa,
            fornecedor=self.fornecedor, numero_nf="NF-987", data_emissao=self.today-timedelta(days=30),
            descricao="Material para obra")
        PagamentoNotaFiscalItem.objects.create(pagamento_nf=self.nf, tipo="produto",
            descricao="Cimento", caixa=self.caixa, valor_total=Decimal("1500"))
        self.boleto = BoletoPagamento.objects.create(pagamento_nf=self.nf, parcela=1,
            numero_doc="BOL-654", vencimento=self.today, valor=Decimal("1500"), status="emitido")

    def context(self, **params):
        request = self.factory.get("/", params)
        request.empresa_ativa = self.empresa
        return _buscar_pagamentos_context(request, self.empresa)

    def test_busca_rapida_por_nome_documento_e_cnpj(self):
        for term in ("Aurora", "NF-987", "BOL-654", "12.345.678/0001-95"):
            with self.subTest(term=term):
                self.assertEqual(len(self.context(busca=term)["resultados"]), 1)

    def test_faixa_de_saldo_e_vencimento_combinados(self):
        context = self.context(busca="Aurora", tipo_valor="valor_a_pagar",
            valor_min="1.000,00", valor_max="2.000,00", vencimento="hoje")
        self.assertFalse(context["form_busca"].errors)
        self.assertEqual(context["totais_busca"]["valor_total"], Decimal("1500"))
        self.assertEqual(len(self.context(busca="Aurora", valor_min="2000,00")["resultados"]), 0)

    def test_tipo_de_data_nao_confunde_emissao_com_vencimento(self):
        params = {"busca": "Aurora", "data_inicio": self.today.isoformat(), "data_fim": self.today.isoformat()}
        self.assertEqual(len(self.context(**params, tipo_data="vencimento")["resultados"]), 1)
        self.assertEqual(len(self.context(**params, tipo_data="emissao")["resultados"]), 0)
        self.assertEqual(len(self.context(**params, tipo_data="pagamento")["resultados"]), 0)

    def test_erro_de_filtro_nao_amplia_busca(self):
        for params in ({"valor_min": "abc"}, {"valor_min": "20", "valor_max": "10"},
                       {"data_inicio": "invalida"}, {"data_inicio": "2026-10-10", "data_fim": "2026-01-01"},
                       {"tipo": "inexistente"}):
            with self.subTest(params=params):
                ctx = self.context(busca="Aurora", **params)
                self.assertTrue(ctx["form_busca"].errors)
                self.assertEqual(ctx["resultados"], [])

    def test_caixa_de_outra_empresa_rejeitado(self):
        other = Empresa.objects.create(razao_social="Outra", cnpj="99999999000199")
        caixa = Caixa.objects.get(empresa=other, tipo=Caixa.Tipo.GERAL)
        ctx = self.context(busca="Aurora", caixa=caixa.pk)
        self.assertTrue(ctx["form_busca"].errors)
        self.assertEqual(ctx["resultados"], [])

    def test_descricao_e_forma(self):
        self.assertEqual(len(self.context(descricao="Cimento", forma="boletos")["resultados"]), 1)
        self.assertEqual(len(self.context(descricao="Cimento", forma="credito")["resultados"]), 0)

    def test_chips_removem_apenas_um_filtro_e_exportacao_descreve_filtros(self):
        ctx = self.context(busca="Aurora", descricao="Cimento", vencimento="hoje")
        self.assertEqual(ctx["filtros_avancados_count"], 3)
        chip = next(c for c in ctx["filtros_chips"] if c["label"].startswith("Descrição"))
        self.assertNotIn("descricao=", chip["url"])
        self.assertIn("busca=Aurora", chip["url"])
        self.assertIn("Vence hoje", _filtros_export_busca_pagamentos(ctx))

    def test_template_renderiza_painel_e_filtros(self):
        from django.contrib.auth.models import AnonymousUser
        request = self.factory.get("/")
        request.empresa_ativa = self.empresa
        request.user = AnonymousUser()
        html = render_to_string("financeiro/busca_pagamentos.html",
            self.context(busca="Aurora", descricao="Cimento"), request=request)
        self.assertIn("Filtros avançados (2)", html)
        self.assertIn('id="filtrosAvancados" class="collapse"', html)
        self.assertIn('name="descricao"', html)
        self.assertIn("Aurora", html)

    def test_imposto_pendente_aparece_sem_ser_tratado_como_pago(self):
        autoridade = AutoridadeTributaria.objects.create(empresa=self.empresa, nome="Receita",
            esfera="federal")
        imposto = PagamentoImposto.objects.create(empresa=self.empresa, caixa=self.caixa,
            autoridade=autoridade, data_vencimento=self.today)
        categoria = CategoriaFinanceira.objects.create(empresa=self.empresa, nome="Impostos",
            tipo="saida", movimentacao_tipo="pag_impostos")
        PagamentoImpostoItem.objects.create(pagamento=imposto, descricao="Guia",
            categoria=categoria, valor_total=Decimal("90"))
        ctx = self.context(busca="Receita")
        self.assertEqual(len(ctx["resultados"]), 1)
        item = ctx["resultados"][0]
        self.assertEqual(item["valor_pago"], Decimal("0"))
        self.assertEqual(item["detalhes_avista"], [])

    def test_bancario_respeita_conta_tipo_e_total(self):
        conta = ContaBancaria.objects.create(empresa=self.empresa, nome="Conta principal",
            banco="Banco teste", conta="123")
        categoria = CategoriaFinanceira.objects.create(empresa=self.empresa, nome="Tarifas",
            tipo="saida", movimentacao_tipo="pag_bancario")
        PagamentoBancarioAvulso.objects.create(empresa=self.empresa, caixa=self.caixa,
            conta_bancaria=conta, categoria=categoria, descricao="Tarifa mensal",
            data_pagamento=self.today, valor=Decimal("30"))
        ctx = self.context(busca="Tarifa", tipo="bancario", conta=conta.pk)
        self.assertEqual(len(ctx["resultados"]), 1)
        self.assertEqual(ctx["totais_busca"]["valor_total"], Decimal("30"))
        self.assertEqual(ctx["totais_busca"]["valor_bancario"], Decimal("30"))
        self.assertEqual(self.context(busca="Tarifa", tipo="nf")["resultados"], [])
        self.assertEqual(self.context(busca="Tarifa", tipo_data="emissao",
            data_inicio=self.today.isoformat())["resultados"], [])

    def test_bancario_parcela_vencida(self):
        from .models import PagamentoBancarioRecorrente, PagamentoBancarioParcela
        conta = ContaBancaria.objects.create(empresa=self.empresa, nome="Conta principal",
            banco="Banco teste", conta="123")
        categoria = CategoriaFinanceira.objects.create(empresa=self.empresa, nome="Tarifas",
            tipo="saida", movimentacao_tipo="pag_bancario")
        recorrencia = PagamentoBancarioRecorrente.objects.create(empresa=self.empresa,
            caixa=self.caixa, conta_bancaria=conta, categoria=categoria, dia_pagamento=1,
            data_inicio=self.today, valor_parcela=Decimal("100"), descricao="Tarifa recorrente")
        PagamentoBancarioParcela.objects.create(recorrencia=recorrencia, numero_parcela=1,
            data_vencimento=self.today-timedelta(days=1), valor=Decimal("100"))
        ctx = self.context(tipo="bancario", vencimento="vencido")
        self.assertEqual(len(ctx["resultados"]), 1)
        self.assertEqual(ctx["resultados"][0]["valor_a_pagar"], Decimal("100"))
        self.assertEqual(self.context(tipo="bancario", status_pago="1")["resultados"], [])

    def test_sem_filtros_nao_carrega_todos_os_pagamentos(self):
        self.assertEqual(self.context()["resultados"], [])

    def test_pdf_e_excel_usam_os_mesmos_resultados(self):
        from .views import buscar_pagamentos_xlsx, buscar_pagamentos_pdf
        from django.contrib.auth import get_user_model
        from io import BytesIO
        from openpyxl import load_workbook
        request = self.factory.get("/", {"busca": "Aurora", "vencimento": "hoje"})
        request.empresa_ativa = self.empresa
        request.user = get_user_model()(username="teste")
        xlsx = buscar_pagamentos_xlsx.__wrapped__(request)
        pdf = buscar_pagamentos_pdf.__wrapped__(request)
        self.assertEqual(xlsx.status_code, 200)
        self.assertEqual(pdf.status_code, 200)
        self.assertTrue(pdf.content.startswith(b"%PDF"))
        wb = load_workbook(BytesIO(xlsx.content))
        text = str(list(wb.active.values))
        self.assertIn("NF-987", text)
        self.assertIn("Vence hoje", text)

    def test_painel_aberto_quando_filtros_invalidos(self):
        from django.contrib.auth.models import AnonymousUser
        request = self.factory.get("/")
        request.empresa_ativa = self.empresa
        request.user = AnonymousUser()
        html = render_to_string("financeiro/busca_pagamentos.html",
            self.context(valor_min="abc"), request=request)
        self.assertIn('id="filtrosAvancados" class="collapse show"', html)


    def test_valor_contem_busca_parte_do_total(self):
        self.assertEqual(len(self.context(valor_contem="500")["resultados"]), 1)
        self.assertEqual(len(self.context(valor_contem="1.500,00")["resultados"]), 1)
        self.assertEqual(self.context(valor_contem="999")["resultados"], [])
        self.assertTrue(self.context(valor_contem="abc")["form_busca"].errors)

    def test_vence_hoje_sozinho_e_com_vencidos(self):
        self.assertEqual(len(self.context(status_hoje="1")["resultados"]), 1)
        self.assertEqual(len(self.context(status_hoje="1", status_vencido="1")["resultados"]), 1)
        self.boleto.vencimento = self.today-timedelta(days=2)
        self.boleto.save()
        self.assertEqual(self.context(status_hoje="1")["resultados"], [])
        self.assertEqual(len(self.context(status_hoje="1", status_vencido="1")["resultados"]), 1)

    def test_campos_principais_nao_contam_como_avancados(self):
        ctx = self.context(q="987", fornecedor="Aurora", valor_contem="500",
            data_inicio="2026-01-01", forma="boletos", status_hoje="1")
        self.assertEqual(ctx["filtros_avancados_count"], 0)

    def test_exportacao_seleciona_ids_e_recalcula_totais(self):
        request = self.factory.get("/", {"busca": "Aurora"})
        request.empresa_ativa = self.empresa
        ctx = _buscar_pagamentos_context(request, self.empresa, selecionados={f"nf:{self.nf.pk}"})
        self.assertEqual(ctx["totais_busca"]["valor_total"], Decimal("1500"))
        ctx = _buscar_pagamentos_context(request, self.empresa, selecionados={"nf:999999"})
        self.assertEqual(ctx["resultados"], [])
        self.assertEqual(ctx["totais_busca"]["valor_total"], Decimal("0"))

    def test_exports_post_excluem_desmarcados(self):
        import json
        from io import BytesIO
        from openpyxl import load_workbook
        from django.contrib.auth import get_user_model
        from .views import buscar_pagamentos_xlsx, buscar_pagamentos_pdf
        segunda = PagamentoNotaFiscal.objects.create(empresa=self.empresa, caixa=self.caixa,
            fornecedor=self.fornecedor, numero_nf="NF-EXCLUIDA")
        PagamentoNotaFiscalItem.objects.create(pagamento_nf=segunda, caixa=self.caixa,
            tipo="produto", descricao="Outro", valor_total=Decimal("99"))
        request = self.factory.post("/?busca=Aurora", {"selecionados": json.dumps([f"nf:{self.nf.pk}"])})
        request.empresa_ativa = self.empresa
        request.user = get_user_model()(username="teste")
        resposta = buscar_pagamentos_xlsx.__wrapped__(request)
        self.assertEqual(resposta.status_code, 200)
        linhas = list(load_workbook(BytesIO(resposta.content)).active.values)
        self.assertIn("NF-987", str(linhas))
        self.assertNotIn("NF-EXCLUIDA", str(linhas))
        self.assertIn("Total Pago", str(linhas))
        self.assertEqual(buscar_pagamentos_pdf.__wrapped__(request).status_code, 200)
        request.POST = request.POST.copy()
        request.POST["selecionados"] = "[]"
        self.assertEqual(buscar_pagamentos_xlsx.__wrapped__(request).status_code, 400)
        request.POST["selecionados"] = '{"invalido": true}'
        self.assertEqual(buscar_pagamentos_pdf.__wrapped__(request).status_code, 400)

    def test_nf_exibe_vencimento_pendente_e_ultima_data_boleto(self):
        from .models import PagamentoNotaFiscalPagamento
        self.boleto.status = "pago"
        self.boleto.valor_pago = self.boleto.valor
        self.boleto.data_pagamento = self.today-timedelta(days=5)
        self.boleto.save()
        BoletoPagamento.objects.create(pagamento_nf=self.nf, parcela=2,
            numero_doc="BOL-2", vencimento=self.today+timedelta(days=10),
            valor=Decimal("100"), status="emitido")
        row = self.context(busca="Aurora")["resultados"][0]
        self.assertEqual(row["vencimento_status"], self.today+timedelta(days=10))
        self.assertEqual(row["data_pagamento"], self.today-timedelta(days=5))
        self.assertEqual(row["forma_pagamento"], "Boletos")
        self.assertEqual(row["forma_detalhe"], "1/2 parcelas pagas")
        PagamentoNotaFiscalPagamento.objects.create(pagamento_nf=self.nf,
            tipo="avista", data=self.today, valor=Decimal("10"))
        row = self.context(busca="Aurora")["resultados"][0]
        self.assertEqual(row["forma_pagamento"], "Misto")
        self.assertEqual(row["data_pagamento"], self.today)

    def test_total_vencido_soma_apenas_saldo_atrasado(self):
        self.boleto.vencimento = self.today-timedelta(days=1)
        self.boleto.valor_pago = Decimal("400")
        self.boleto.save()
        for numero, status, vencimento in (
            (2, "emitido", self.today), (3, "emitido", self.today+timedelta(days=2)),
            (4, "pago", self.today-timedelta(days=4)),
            (5, "cancelado", self.today-timedelta(days=4))):
            BoletoPagamento.objects.create(pagamento_nf=self.nf, parcela=numero, numero_doc=f"TESTE-{numero}",
                valor=Decimal("50"), status=status, vencimento=vencimento)
        ctx = self.context(busca="Aurora")
        self.assertEqual(ctx["resultados"][0]["valor_vencido"], Decimal("1100"))
        self.assertEqual(ctx["totais_busca"]["valor_vencido"], Decimal("1100"))

    def test_vence_hoje_nao_compoe_total_vencido(self):
        ctx = self.context(busca="Aurora")
        self.assertEqual(ctx["totais_busca"]["valor_vencido"], Decimal("0"))

    def test_totais_sao_rodape_da_tabela_principal(self):
        from html.parser import HTMLParser
        from django.contrib.auth.models import AnonymousUser

        class TabelasParser(HTMLParser):
            def __init__(self):
                super().__init__()
                self.tables = []
                self.footer_tables = []

            def handle_starttag(self, tag, attrs):
                if tag == "table":
                    self.tables.append(dict(attrs).get("id"))
                if tag == "tfoot":
                    self.footer_tables.append(self.tables[-1] if self.tables else None)

            def handle_endtag(self, tag):
                if tag == "table":
                    self.tables.pop()

        request = self.factory.get("/")
        request.empresa_ativa = self.empresa
        request.user = AnonymousUser()
        html = render_to_string("financeiro/busca_pagamentos.html",
            self.context(busca="Aurora"), request=request)
        parser = TabelasParser()
        parser.feed(html)
        self.assertEqual(parser.footer_tables, ["financeiroResultadosTable"])
        self.assertLess(html.index('class="fin-totals-footer"'), html.index('aria-label="Filtros aplicados"'))

class ApresentacaoResultadosTests(TestCase):
    def item(self, emissao, vencimentos, pagamentos=()):
        return {"kind": "nf", "entidade": "Fornecedor", "numero_doc": "NF-10",
                "_datas": {"emissao": [emissao], "vencimento": list(vencimentos), "pagamento": list(pagamentos)},
                "_pendentes": list(vencimentos), "valor_total": Decimal("1000"),
                "valor_pago": Decimal("400"), "valor_a_pagar": Decimal("600"),
                "detalhes_boletos": [{"doc": "BOL-1", "vencimento": d} for d in vencimentos]}

    def test_vencidos_priorizam_saldo_e_vencimento_real(self):
        from .busca import preparar_resultados
        hoje = timezone.localdate()
        older = self.item(hoje, [hoje-timedelta(days=20), hoje-timedelta(days=5)])
        newer = self.item(hoje, [hoje-timedelta(days=2)])
        rows = [newer, older]
        layout = preparar_resultados(rows, {"vencimento": "vencido"})
        self.assertIs(rows[0], older)
        self.assertEqual(older["data_destaque"], hoje-timedelta(days=20))
        self.assertEqual(older["valor_destaque"], Decimal("600"))
        self.assertEqual(layout["valor_label"], "Saldo pendente")
        self.assertTrue(older["detalhes_boletos"][0]["destaque_contexto"])

    def test_data_exibida_corresponde_ao_periodo_e_tipo(self):
        from .busca import preparar_resultados
        hoje = timezone.localdate()
        row = self.item(hoje-timedelta(days=30), [hoje-timedelta(days=5), hoje])
        preparar_resultados([row], {"tipo_data": "vencimento", "data_inicio": hoje, "data_fim": hoje})
        self.assertEqual(row["data_destaque"], hoje)

    def test_faixa_total_nao_destaca_saldo(self):
        from .busca import preparar_resultados
        hoje = timezone.localdate()
        row = self.item(hoje, [hoje])
        preparar_resultados([row], {"status_aberto": True, "valor_min": Decimal("500")})
        self.assertEqual(row["valor_destaque"], Decimal("1000"))

    def test_pago_prioriza_ultimo_pagamento(self):
        from .busca import preparar_resultados
        hoje = timezone.localdate()
        row = self.item(hoje-timedelta(days=30), [], [hoje-timedelta(days=2), hoje])
        preparar_resultados([row], {"status_pago": True})
        self.assertEqual(row["data_destaque"], hoje)
        self.assertEqual(row["valor_destaque"], Decimal("400"))
        self.assertEqual(row["datas_adicionais"], 1)

    def test_valores_e_composicao_preservados(self):
        from .busca import preparar_resultados
        hoje = timezone.localdate()
        row = self.item(hoje, [hoje])
        preparar_resultados([row], {"tipo_valor": "valor_pago", "busca": "NF-10"})
        self.assertEqual(row["valor_total"], Decimal("1000"))
        self.assertEqual(row["valor_a_pagar"], Decimal("600"))
        self.assertTrue(row["destaque_documento"])
